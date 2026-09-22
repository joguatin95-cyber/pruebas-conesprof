"""Adaptadores de lectura de archivos tabulares.

Patrones aplicados:

- **Strategy**: cada formato es una implementacion intercambiable de
  LectorTabular. Agregar ODS manana no obliga a tocar el caso de uso.
- **Factory Method**: `lector_para()` elige la estrategia segun la extension, de
  modo que quien llama no necesita conocer las clases concretas.
"""

import csv
import io
import os

from app.aplicacion.puertos import ErrorDeLectura, LectorTabular, TablaLeida

EXTENSIONES_EXCEL = (".xlsx", ".xlsm")
EXTENSIONES_CSV = (".csv", ".txt")
EXTENSIONES_ADMITIDAS = EXTENSIONES_EXCEL + EXTENSIONES_CSV

CODIFICACIONES = ("utf-8-sig", "utf-8", "cp1252", "latin-1")
SEPARADORES = ";,\t|"


class LectorExcel(LectorTabular):
    def leer(self, contenido: bytes, hoja: str | None = None) -> TablaLeida:
        from openpyxl import load_workbook

        try:
            libro = load_workbook(io.BytesIO(contenido), data_only=True, read_only=True)
        except Exception as error:  # openpyxl lanza tipos muy variados
            raise ErrorDeLectura(
                "No se pudo abrir el archivo de Excel. Verifique que sea un .xlsx "
                "valido y no este protegido con contrasena."
            ) from error

        try:
            if hoja:
                if hoja not in libro.sheetnames:
                    raise ErrorDeLectura(
                        "La hoja '%s' no existe. Hojas disponibles: %s."
                        % (hoja, ", ".join(libro.sheetnames))
                    )
                pagina = libro[hoja]
            else:
                pagina = libro[libro.sheetnames[0]]

            filas = list(pagina.iter_rows(values_only=True))
            titulo = pagina.title
        finally:
            libro.close()

        if not filas:
            raise ErrorDeLectura("La hoja del archivo esta vacia.")
        return TablaLeida(list(filas[0]), [list(f) for f in filas[1:]],
                          "Hoja: %s" % titulo)


class LectorCsv(LectorTabular):
    def leer(self, contenido: bytes, hoja: str | None = None) -> TablaLeida:
        texto = self._decodificar(contenido)
        separador = self._detectar_separador(texto[:8000])
        filas = list(csv.reader(texto.splitlines(), delimiter=separador))
        if not filas:
            raise ErrorDeLectura("El archivo esta vacio.")
        return TablaLeida(filas[0], filas[1:],
                          "CSV (separador '%s')" % separador)

    @staticmethod
    def _decodificar(contenido: bytes) -> str:
        for codificacion in CODIFICACIONES:
            try:
                return contenido.decode(codificacion)
            except UnicodeDecodeError:
                continue
        raise ErrorDeLectura(
            "No se pudo leer el archivo: la codificacion no es reconocible. "
            "Guardelo de nuevo como CSV UTF-8."
        )

    @staticmethod
    def _detectar_separador(muestra: str) -> str:
        try:
            return csv.Sniffer().sniff(muestra, delimiters=SEPARADORES).delimiter
        except csv.Error:
            # Si no se puede detectar, se elige el separador mas frecuente.
            return max(SEPARADORES, key=muestra.count)


def lector_para(nombre_archivo: str) -> LectorTabular:
    """Factory: devuelve la estrategia de lectura que corresponde al archivo."""
    extension = os.path.splitext(nombre_archivo or "")[1].lower()
    if extension in EXTENSIONES_EXCEL:
        return LectorExcel()
    if extension in EXTENSIONES_CSV:
        return LectorCsv()
    raise ErrorDeLectura(
        "Formato no admitido (%s). Use %s."
        % (extension or "sin extension", ", ".join(EXTENSIONES_ADMITIDAS))
    )


def hojas_de(contenido: bytes) -> list[str]:
    """Nombres de las hojas de un Excel, para ofrecerlas en la pantalla web."""
    from openpyxl import load_workbook

    try:
        libro = load_workbook(io.BytesIO(contenido), read_only=True)
    except Exception:  # noqa: BLE001 - si no se puede abrir, no hay hojas que ofrecer
        return []
    try:
        return list(libro.sheetnames)
    finally:
        libro.close()
